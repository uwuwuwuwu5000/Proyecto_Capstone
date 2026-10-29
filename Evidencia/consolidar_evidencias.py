"""Reúne las tres corridas automáticas en un solo CSV y un solo informe HTML.

  - Portal web (rules-tests, con el emulador de Firestore)   -> portal-reglas.xml
  - Paquete comun (lógica compartida)                         -> comun.xml
  - Pruebas de pantalla con Playwright (emuladores Auth+DB)   -> e2e/corrida-final/e2e-portal.xml

Uso:  python consolidar_evidencias.py
Crea: Evidencias_pruebas_automaticas.csv, Informe_pruebas_automaticas.html y plan_estado.json
"""
import csv, html, json, re, sys, xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
FECHA = datetime.now().strftime('%d-%m-%Y')

# Se reutilizan las reglas de asignación de IDs de asignar_ids.py sin ejecutar su parte de salida.
fuente = (AQUI / 'asignar_ids.py').read_text(encoding='utf8')
exec(compile(fuente[: fuente.index('ruta = Path(sys.argv[1])')], 'asignar_ids', 'exec'))  # define REGLAS e ids_de

ORIGENES = [
    ('Portal web (rules-tests)', AQUI / 'portal-reglas.xml'),
    ('Paquete comun', AQUI / 'comun.xml'),
    ('Playwright (pantallas del portal)', AQUI / 'e2e' / 'corrida-final' / 'e2e-portal.xml'),
]

# Pruebas que el plan separó por plataforma (ID-W para el portal, ID-M para la app): la evidencia de este
# repositorio es del portal, así que se asigna a la fila -W.
def _ids_separados():
    try:
        import openpyxl
        hoja = openpyxl.load_workbook(RAIZ / 'Plan_de_Pruebas_Ciudad_Alerta_2026_limpio.xlsx', read_only=True)['Plan de prueba']
        return {r[2][:-2] for r in hoja.iter_rows(min_row=2, values_only=True) if r[2] and str(r[2]).endswith('-W')}
    except Exception:
        return set()


SEPARADOS = _ids_separados()


def a_portal(ids):
    return ','.join((i + '-W') if i in SEPARADOS else i for i in filter(None, ids.split(',')))


PREFIJO = {'Portal web (rules-tests)': 'portal-web/rules-tests/', 'Paquete comun': 'comun/test/', 'Playwright (pantallas del portal)': 'portal-web/e2e/'}
filas = []
for origen, ruta in ORIGENES:
    if not ruta.exists():
        print('FALTA', ruta)
        continue
    for caso in ET.parse(ruta).getroot().iter('testcase'):
        archivo = caso.get('classname', '').split('/')[-1].split('\\')[-1]
        nombre = caso.get('name', '')
        fallo = caso.find('failure') is not None or caso.find('error') is not None
        omitida = caso.find('skipped') is not None
        estado = 'FALLÓ' if fallo else ('OMITIDA' if omitida else 'OK')
        filas.append([origen, a_portal(ids_de(archivo, nombre)), archivo, nombre, estado, caso.get('time', ''), FECHA])

with open(AQUI / 'Evidencias_pruebas_automaticas.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['Origen', 'IDs del plan', 'Archivo de prueba', 'Prueba', 'Resultado', 'Segundos', 'Fecha de ejecución'])
    w.writerows(filas)

# ---- Estado por ID del plan ----
por_id = defaultdict(lambda: {'ok': 0, 'mal': 0, 'origenes': set(), 'archivos': set()})
for origen, ids, archivo, nombre, estado, _, _ in filas:
    for i in filter(None, ids.split(',')):
        d = por_id[i]
        d['ok' if estado == 'OK' else 'mal'] += 1
        d['origenes'].add(origen)
        d['archivos'].add(PREFIJO[origen] + archivo)
salida = {i: {'ok': d['ok'], 'mal': d['mal'], 'origenes': sorted(d['origenes']), 'archivos': sorted(d['archivos'])} for i, d in por_id.items()}
(AQUI / 'plan_estado.json').write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding='utf8')

# ---- Informe HTML único ----
try:
    import openpyxl
    plan = openpyxl.load_workbook(RAIZ / 'Plan_de_Pruebas_Ciudad_Alerta_2026_limpio.xlsx', read_only=True)['Plan de prueba']
    nombres = {r[2]: r[3] for r in plan.iter_rows(min_row=2, values_only=True) if r[2]}
except Exception:  # sin el Excel el informe sale igual, solo sin nombres
    nombres = {}

color = {'OK': '#d4edda', 'FALLÓ': '#f8d7da', 'OMITIDA': '#fff3cd'}
resumen = []
for origen, _ in ORIGENES:
    propias = [f for f in filas if f[0] == origen]
    ok = sum(f[4] == 'OK' for f in propias)
    resumen.append((origen, len(propias), ok, len(propias) - ok))
total = len(filas)
total_ok = sum(f[4] == 'OK' for f in filas)

filas_resumen = ''.join(f'<tr><td>{html.escape(o)}</td><td>{n}</td><td>{ok}</td><td>{mal}</td></tr>' for o, n, ok, mal in resumen)
filas_ids = ''.join(
    f'<tr style="background:{"#d4edda" if d["mal"] == 0 else "#f8d7da"}"><td>{html.escape(i)}</td><td>{html.escape(str(nombres.get(i, "")))}</td>'
    f'<td>{d["ok"]}</td><td>{d["mal"]}</td><td>{html.escape(", ".join(d["origenes"]))}</td></tr>'
    for i, d in sorted(salida.items())
)
filas_detalle = ''.join(
    f'<tr style="background:{color[e]}"><td>{html.escape(o)}</td><td>{html.escape(i)}</td><td>{html.escape(n)}</td><td>{e}</td><td>{t}</td></tr>'
    for o, i, a, n, e, t, _ in filas
)
fallidas = [f for f in filas if f[4] != 'OK']
bloque_fallos = (
    '<p class="ok">Ninguna prueba falló.</p>'
    if not fallidas
    else '<ul>' + ''.join(f'<li>{html.escape(f[0])}: {html.escape(f[3])}</li>' for f in fallidas) + '</ul>'
)

(AQUI / 'Informe_pruebas_automaticas.html').write_text(
    f'''<!doctype html><meta charset="utf-8"><title>Informe de pruebas automáticas · Ciudad Alerta</title>
<style>body{{font:14px sans-serif;margin:24px;max-width:1200px}}table{{border-collapse:collapse;width:100%;margin-bottom:24px}}
td,th{{border:1px solid #ccc;padding:4px 8px;text-align:left;vertical-align:top}}th{{background:#eee}}.ok{{color:#155724;font-weight:bold}}
h2{{margin-top:28px}}</style>
<h1>Informe de pruebas automáticas · Ciudad Alerta</h1>
<p>Fecha de ejecución: <b>{FECHA}</b>. Resultado general: <b>{total_ok} de {total} pruebas OK</b>.</p>
<p>Un solo informe con las tres corridas: reglas de seguridad del portal (emulador de Firestore), lógica compartida (paquete <i>comun</i>) y
pantallas del portal con Playwright (emuladores de Auth y Firestore: no se toca la base de datos real). El detalle por prueba está en
<code>Evidencias_pruebas_automaticas.csv</code>.</p>
<h2>Resumen por origen</h2>
<table><tr><th>Origen</th><th>Pruebas</th><th>OK</th><th>Con fallo</th></tr>{filas_resumen}</table>
<h2>Pruebas con fallo u omitidas</h2>{bloque_fallos}
<h2>Estado por ID del plan de pruebas ({len(salida)} IDs con evidencia)</h2>
<table><tr><th>ID</th><th>Prueba del plan</th><th>OK</th><th>Fallos</th><th>Origen de la evidencia</th></tr>{filas_ids}</table>
<h2>Detalle de cada prueba automática</h2>
<table><tr><th>Origen</th><th>IDs del plan</th><th>Prueba</th><th>Resultado</th><th>Seg.</th></tr>{filas_detalle}</table>''',
    encoding='utf-8',
)
print(f'{total_ok} OK de {total}; {len(salida)} IDs del plan con evidencia; {len(fallidas)} con fallo u omitidas')
for o, n, ok, mal in resumen:
    print(f'  {o}: {ok}/{n}')
